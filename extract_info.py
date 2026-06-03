import docx
import openpyxl

print("--- Timetable Constraints.docx ---")
try:
    doc = docx.Document('Timetable Constraints.docx')
    for p in doc.paragraphs:
        if p.text.strip():
            print(p.text)
except Exception as e:
    print(f"Error reading docx: {e}")

print("\n--- Timetable Hours Distribution ---")
try:
    wb = openpyxl.load_workbook('Timetable Hours Distribution 25-26 Even Sem.xlsx', data_only=True)
    sheet = wb.active
    for row in sheet.iter_rows(values_only=True):
        print(row)
except Exception as e:
    print(f"Error reading xlsx: {e}")
